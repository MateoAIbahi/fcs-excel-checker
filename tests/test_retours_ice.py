"""
Tests de non-regression des retours ICE (septembre 2026).
Donnees synthetiques reproduisant les captures transmises par le client :
aucun fichier client n'est necessaire.

    python -m pytest tests -q
"""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.excel_parser import (
    detect_cal_mode,
    extract_cal_sheet,
    is_retained_choice,
    parse_sheet_frames,
    resolve_function,
)
from src.comparator import (
    compare_all,
    compare_section,
    COLUMN_OK,
    COLUMN_EXCEL_ONLY,
    COLUMN_FCS_ONLY,
    STATUS_EXCEL_ONLY,
    STATUS_FCS_ONLY,
    STATUS_OK,
    STATUS_OK_CAL,
    STATUS_OK_EQUIVALENCE,
)
from src.report_generator import generate_excel_report
from src.utils import is_negative_answer


# --------------------------------------------------------------------------
# Onglet CAL indice K, avec bloc migration IF-TG en fin d'onglet
# --------------------------------------------------------------------------

def cal_frame_k():
    rows = [
        ["Code", "Désignation", "", "Choix des fonctions", "Séléction Choix = C"],
        [None, "Fonctions numérisées dans le CCN", None, None, None],
        ["TG", "Système traitant des fonctions numérisées", None, "Base", None],
        ["TGSI", "Système traitant des fonctions numérisées sécurisées", None, "Base", None],
        ["SYNO", "Synoptique", None, "Base", None],
        ["ALTECH", "Alarme Incendie bâtiment uniquement", None, "Base", "non suite à la FQR 04"],
        ["CA", "Contrôle d'Accès", None, "Option", "non"],
        ["SIRSUTCTCO2", "Gestion de la Sirène commune SUTCT et CO2", None, "Option", "Oui TCT"],
        ["SIRSF6", "Gestion de la Sirène SF6", None, "Option", "non (car que sur poste blindé)"],
        ["BLOCREGT", "Fonction Blocage des Régleurs", None, "Option", "non (car TR Enedis)"],
        ["TELEAL", "Téléalarme", None, "Choix", "Téléalarme"],
        [None, "Fonctions ou équipements interfacés", None, None, None],
        ["IF-TG", "Fonctions ou équipements du CCN pour Migration (pièce T9)",
         "Non utilisé pour un poste neuf", None, None],
        [None, 'Calculateur de migration T3 "CAL IF-TG" 48V', "Pour PMP", "O", "X"],
        [None, 'Calculateur de migration T3 "CAL IF-TG" 127V', "Pour liaison", "O", None],
    ]
    return pd.DataFrame(rows)


def cal_frame_h():
    rows = [
        ["Code", "Désignation", "", "Choix", "Options retenues par DI"],
        [None, "Fonctions numérisées dans le CCN", None, None, None],
        ["SYNO", "Synoptique", None, "Base", "X"],
        ["CA", "Contrôle d'Accès", None, "Option", "X"],
        ["SIRSF6", "Gestion de la Sirène SF6", None, "Option", None],
        [None, "Fonctions ou équipements interfacés", None, None, None],
    ]
    return pd.DataFrame(rows)


def test_negative_answer():
    assert is_negative_answer("non (car que sur poste blindé)")
    assert is_negative_answer("Non suite à la FQR 04")
    assert is_negative_answer("NON")
    assert not is_negative_answer("Oui TCT")
    assert not is_negative_answer("Néant à signaler")
    assert not is_negative_answer("Téléalarme")
    assert not is_negative_answer(None)


def test_x_du_bloc_migration_ne_bascule_pas_en_mode_marqueur():
    assert detect_cal_mode(cal_frame_k()) == "decision"


def test_mode_marqueur_toujours_detecte_en_indice_h():
    functions, _, mode, _ = extract_cal_sheet(cal_frame_h())
    assert mode == "marqueur"
    assert functions == {"SYNO", "CA"}


def test_selection_cal_indice_k():
    functions, _, _, all_codes = extract_cal_sheet(cal_frame_k())
    assert functions == {"SYNO", "SIRSUTCTCO2", "TELEAL"}
    assert "IF-TG" in all_codes


def test_base_toujours_retenue_sans_refus():
    assert is_retained_choice("Base", None)
    assert is_retained_choice("Base", "C")
    assert not is_retained_choice("Option", None)


def tg_fcs(functions):
    return {"tranches": {"T.GENE": {
        "LibelléLongTranche": "TRANCHE GENERALE", "LibelléTT": "",
        "CodeSchémathèqueTT": "",
        "FonctionsNumériséesCCN": functions, "EquipementsTiers": {},
    }}}


def test_tg_tgsi_ignores_et_if_tg_trouve():
    parsed = parse_sheet_frames({"PdG ": pd.DataFrame(), "CAL": cal_frame_k()})
    fcs = tg_fcs({"SYNO": "Synoptique", "TG": "Système", "IF-TG": "Migration"})
    comparison = compare_all(
        fcs, {"tg.xls": {"tranche": "T.GENE", "method": "tranche_generale"}},
        {"tg.xls": parsed},
    )
    rows = comparison["tranches"]["T.GENE"]["sections"]["FonctionsNumériséesCCN"]
    by_code = {r["Fonction FCS"] or r["Fonction Excel"]: r["Statut"] for r in rows}

    assert "TG" not in by_code and "TGSI" not in by_code
    assert by_code["SYNO"] == STATUS_OK
    assert by_code["IF-TG"] == STATUS_OK_CAL
    assert by_code["SIRSUTCTCO2"] == STATUS_EXCEL_ONLY
    assert "SIRSF6" not in by_code and "ALTECH" not in by_code


def test_if_tg_absent_du_fcs_non_signale():
    parsed = parse_sheet_frames({"PdG ": pd.DataFrame(), "CAL": cal_frame_k()})
    rows = compare_section(parsed, {"SYNO": ""}, "FonctionsNumériséesCCN")
    assert all(r["Fonction Excel"] != "IF-TG" for r in rows)


# --------------------------------------------------------------------------
# Equivalence UT (FCS) <-> PDBN (nomenclature)
# --------------------------------------------------------------------------

def bt_frame(mnemonic):
    return pd.DataFrame([
        ["Mnémonique", "Désignation", "Type option"],
        [mnemonic, "Protection différentielle de barres", "Base"],
    ])


def test_ut_pdbn_equivalence():
    parsed = parse_sheet_frames({"2-Basse Tension": bt_frame("PDBN")})
    found, method, evidence = resolve_function(
        parsed, "EquipementsTiers", "UT",
        "Unité de travée différentielle de barres numérique",
    )
    assert (found, method, evidence) == (True, "equivalence", "PDBN")

    rows = compare_section(
        parsed,
        {"UT": "Unite de travee differentielle de barres numerique"},
        "EquipementsTiers",
    )
    assert [r["Statut"] for r in rows] == [STATUS_OK_EQUIVALENCE]


def test_ut_autre_libelle_non_rapproche():
    parsed = parse_sheet_frames({"2-Basse Tension": bt_frame("PDBN")})
    found, method, _ = resolve_function(
        parsed, "EquipementsTiers", "UT", "Unité de traitement",
    )
    assert method != "equivalence"


def test_ut_pdbn_indice():
    parsed = parse_sheet_frames({"2-Basse Tension": bt_frame("PDBN1")})
    found, method, evidence = resolve_function(
        parsed, "EquipementsTiers", "UT",
        "Unité de travée différentielle de barres numérique",
    )
    assert (method, evidence) == ("equivalence", "PDBN1")


# --------------------------------------------------------------------------
# Rapport
# --------------------------------------------------------------------------

def test_titres_colonnes_resume():
    parsed = parse_sheet_frames({"PdG ": pd.DataFrame(), "CAL": cal_frame_k()})
    comparison = compare_all(
        tg_fcs({"SYNO": "", "ABSENT": ""}),
        {"tg.xls": {"tranche": "T.GENE", "method": "tranche_generale"}},
        {"tg.xls": parsed},
    )
    report = generate_excel_report(comparison)
    sheet = load_workbook(BytesIO(report.getvalue()))["Résumé"]
    headers = [c.value for c in sheet[1]]
    assert COLUMN_OK in headers
    assert COLUMN_EXCEL_ONLY in headers
    assert COLUMN_FCS_ONLY in headers
    assert "OK" not in headers

    values = dict(zip(headers, [c.value for c in sheet[2]]))
    assert values[COLUMN_OK] == 1
    assert values[COLUMN_FCS_ONLY] == 1
    assert values[COLUMN_EXCEL_ONLY] == 2      # SIRSUTCTCO2, TELEAL

    details = load_workbook(BytesIO(report.getvalue()))["Détails"]
    statuses = {row[5] for row in details.iter_rows(min_row=2, values_only=True)}
    assert STATUS_FCS_ONLY in statuses


def test_mnemonique_indice_deterministe():
    parsed = {"EquipementsTiers": set(),
              "mnemonics": {"DDS10", "DDS2", "DDS1"}, "labels": []}
    for _ in range(20):
        assert resolve_function(parsed, "EquipementsTiers", "DDS")[2] == "DDS1"
